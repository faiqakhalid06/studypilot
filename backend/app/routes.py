from io import BytesIO

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai_quiz import AiQuizRequest, make_ai_quiz
from app.tutor import TutorAnswer, TutorQuestion, ask_tutor

from app.database import get_db
from app.models import Subject, Topic
from app.planner import make_plan
from app.progress import find_topics
from app.quiz import make_quiz
from app.schemas import (
    AnswerIn, PlanRequest, PlanResponse, QuizRequest, QuizResponse, SubjectIn,
    SubjectOut, TopicIn, TopicOut, TopicsRequest, TopicsResponse,
)

router = APIRouter(prefix="/api")


def topic_out(t: Topic) -> TopicOut:
    return TopicOut(id=t.id, name=t.name, notes=t.notes, right=t.right_answers, total=t.total_answers)


def subject_out(s: Subject) -> SubjectOut:
    return SubjectOut(id=s.id, name=s.name, exam_date=s.exam_date, topics=[topic_out(t) for t in s.topics])


def subject_in(s: Subject) -> SubjectIn:
    return SubjectIn(
        name=s.name,
        exam_date=s.exam_date,
        topics=[TopicIn(name=t.name, notes=t.notes, right=t.right_answers, total=t.total_answers) for t in s.topics],
    )


def saved_subjects(db: Session) -> list[SubjectIn]:
    rows = db.scalars(select(Subject).order_by(Subject.id)).all()
    return [subject_in(s) for s in rows]


@router.post("/subjects", response_model=SubjectOut, status_code=201)
def add_subject(body: SubjectIn, db: Session = Depends(get_db)):
    subject = Subject(
        name=body.name,
        exam_date=body.exam_date,
        topics=[
            Topic(name=t.name, notes=t.notes, right_answers=t.right, total_answers=t.total)
            for t in body.topics
        ],
    )
    db.add(subject)
    db.commit()
    return subject_out(subject)


@router.get("/subjects", response_model=list[SubjectOut])
def list_subjects(db: Session = Depends(get_db)):
    return [subject_out(s) for s in db.scalars(select(Subject).order_by(Subject.id)).all()]


@router.delete("/subjects/{subject_id}", status_code=204)
def delete_subject(subject_id: int, db: Session = Depends(get_db)):
    subject = db.get(Subject, subject_id)
    if subject is None:
        raise HTTPException(status_code=404, detail="Subject not found")
    db.delete(subject)
    db.commit()
    return Response(status_code=204)


@router.post("/topics/{topic_id}/answers", response_model=TopicOut)
def record_answer(topic_id: int, body: AnswerIn, db: Session = Depends(get_db)):
    topic = db.get(Topic, topic_id)
    if topic is None:
        raise HTTPException(status_code=404, detail="Topic not found")
    topic.total_answers += 1
    if body.correct:
        topic.right_answers += 1
    db.commit()
    return topic_out(topic)


@router.get("/plan", response_model=PlanResponse)
def saved_plan(
    hours_per_day: float = Query(2, ge=0.5, le=12),
    days: int = Query(7, ge=1, le=14),
    db: Session = Depends(get_db),
):
    return make_plan(PlanRequest(subjects=saved_subjects(db), hours_per_day=hours_per_day, days=days))


@router.get("/quiz", response_model=QuizResponse)
def saved_quiz(
    count: int = Query(5, ge=1, le=20),
    seed: int | None = None,
    db: Session = Depends(get_db),
):
    return make_quiz(QuizRequest(subjects=saved_subjects(db), count=count, seed=seed))


@router.get("/weak-topics", response_model=TopicsResponse)
def saved_weak_topics(db: Session = Depends(get_db)):
    return find_topics(TopicsRequest(subjects=saved_subjects(db)).subjects)

@router.post("/tutor/ask", response_model=TutorAnswer)
def tutor_ask(body: TutorQuestion, db: Session = Depends(get_db)):
    return ask_tutor(db, body)

@router.post("/quiz/ai", response_model=QuizResponse)
def ai_quiz(body: AiQuizRequest, db: Session = Depends(get_db)):
    return make_ai_quiz(QuizRequest(subjects=saved_subjects(db), count=body.count), body.asked)

def with_periods(text: str) -> str:
    """Slides and Word bullets have no full stops, so add them. Search and quiz work on sentences."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    return "\n".join(l if l[-1] in ".!?:" else l + "." for l in lines)


@router.post("/extract-text")
async def extract_text(file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File is too big (max 20 MB)")
    name = (file.filename or "").lower()
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif name.endswith(".docx"):
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.extend(c.text for c in row.cells)
            text = with_periods("\n".join(parts))
        elif name.endswith(".pptx"):
            parts = []
            for slide in Presentation(BytesIO(data)).slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text_frame.text)
            text = with_periods("\n".join(parts))
        elif name.endswith((".txt", ".md")):
            text = data.decode("utf-8", errors="ignore")
        else:
            raise HTTPException(status_code=400, detail="Please upload a PDF, Word (.docx), PowerPoint (.pptx) or TXT file")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this file")
    text = " ".join(text.split())
    if not text:
        raise HTTPException(
            status_code=400,
            detail="No text found. This may be a scanned PDF (images), which is not supported yet.",
        )
    return {"text": text[:200000]}